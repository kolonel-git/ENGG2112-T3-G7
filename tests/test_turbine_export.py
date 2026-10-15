"""Checks for turbine separation, exact-time joins, and conflicting sensor records."""

import csv
import os
import subprocess
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from wt_predmaint.data.clean_workbook import clean_workbook
from wt_predmaint.data.turbine_export import (
    FEATURES,
    HEADERS,
    NS,
    export_turbines,
    publish_workbook,
    timestamp,
)


class TurbineExportTests(unittest.TestCase):
    def test_day_first_pressure_dates(self) -> None:
        self.assertEqual(timestamp("1/10/2018 0:00"), "2018-10-01 00:00:00")
        self.assertEqual(timestamp("13/10/2018 12:30:15"), "2018-10-13 12:30:15")
        self.assertEqual(timestamp("2018-06-01 00:00:00"), "2018-06-01 00:00:00")

    def test_rounded_boundary_duplicates_retain_precision(self) -> None:
        self.write("tblSCTurPress", ["TimeStamp", "StationId", "wtc_GearPres_mean"], [
            ["2018-06-01 00:00:00", "2304510", "2.973145008"],
            ["2018-06-01 00:00:00", "2304510", "2.973145008087158"],
        ])
        report = export_turbines(self.root, self.output)
        self.assertEqual(report["rounded_duplicates_merged"]["tblSCTurPress"], 1)
        workbook = load_workbook(self.output, read_only=True)
        try:
            values = list(workbook["Station_2304510"].values)
            self.assertEqual(values[1][2], 2.973145008087158)
        finally:
            workbook.close()

    def test_multiple_months_share_turbine_sheets(self) -> None:
        for table, columns in FEATURES.items():
            with (self.root / f"{table}_october.csv").open("w", newline="") as file:
                writer = csv.writer(file)
                writer.writerow(["TimeStamp", "StationId", columns[0]])
                time = "1/10/2018 0:00" if table == "tblSCTurPress" else "2018-10-01 00:00:00"
                writer.writerow([time, "2304510", "9"])
        report = export_turbines(self.root, self.output)
        self.assertEqual(report["rows_per_station"]["2304510"], 3)
        workbook = load_workbook(self.output, read_only=True)
        try:
            values = list(workbook["Station_2304510"].values)
            self.assertEqual(values[-1][0], "2018-10-01 00:00:00")
            self.assertEqual(values[-1][2], 9)
        finally:
            workbook.close()

    def setUp(self) -> None:
        self.work = tempfile.TemporaryDirectory()
        self.addCleanup(self.work.cleanup)
        self.root = Path(self.work.name)
        self.output = self.root / "output.xlsx"
        for table, columns in FEATURES.items():
            self.write(table, ["TimeStamp", "StationId", columns[0]], [
                ["2018-06-01 00:00:00", "2304510", "1.5"],
                ["2018-06-01 00:00:00", "2304511", "2.5"],
            ])
        self.write("tblAlarmLog", ["TimeOn", "TimeOff", "StationNr", "Alarmcode"], [
            ["2018-06-01 00:00:31", "", "2304510", "25"],
            ["2018-06-01 00:00:31", "2018-06-01 00:01:00", "2304510", "20"],
            ["2018-06-01 00:00:00", "", "2304511", "99"],
        ])

    def write(self, table: str, headers: list[str], rows: list[list[str]]) -> None:
        with (self.root / f"{table}_test.csv").open("w", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(headers)
            writer.writerows(rows)

    def test_exact_time_merge_and_separate_stations(self) -> None:
        report = export_turbines(self.root, self.output)
        self.assertEqual(report["rows_per_station"], {"2304510": 2, "2304511": 1})
        self.assertEqual(report["alarm_only_rows"], 1)
        self.assertIn("wtc_InvUpTmp_mean", report["missing_features"]["tblSCTurTemp"])
        with ZipFile(self.output) as workbook:
            self.assertIsNone(workbook.testzip())
            for path in workbook.namelist():
                ET.fromstring(workbook.read(path))
            metadata = ET.fromstring(workbook.read("xl/workbook.xml"))
            self.assertEqual([s.attrib["name"] for s in metadata.findall(f"{{{NS}}}sheets/*")],
                             ["Station_2304510", "Station_2304511"])
            sheet = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
            rows = sheet.findall(f"{{{NS}}}sheetData/{{{NS}}}row")
            alarm_row = rows[2]
            text = " ".join(alarm_row.itertext())
            self.assertIn("2018-06-01 00:00:31", text)
            self.assertIn("25; 20", text)
            self.assertIn("2018-06-01 00:01:00", text)
            self.assertEqual(len(alarm_row), 4)  # time, ID, codes, event details; no sensors
            sensor_row = rows[1]
            self.assertEqual(len(sensor_row), 5)  # time, ID, one feature from each table
            self.assertEqual(len(rows[0]), len(HEADERS))

    def test_conflicting_duplicates_fail(self) -> None:
        self.write("tblSCTurPress", ["TimeStamp", "StationId", "wtc_GearPres_mean"], [
            ["2018-06-01 00:00:00", "2304510", "1"],
            ["2018-06-01 00:00:00", "2304510", "2"],
        ])
        with self.assertRaisesRegex(ValueError, "Conflicting duplicate"):
            export_turbines(self.root, self.output)
        self.assertFalse(self.output.exists())

    def test_cleaning_keeps_complete_sensors_without_alarms(self) -> None:
        export_turbines(self.root, self.output)
        processed = self.root / "processed.xlsx"
        report = clean_workbook(self.output, processed)
        station = report["sheets"]["Station_2304510"]
        self.assertEqual(station["kept_rows"], 1)
        self.assertEqual(station["removed_rows"], 1)
        self.assertEqual(station["alarm_events_in_removed_rows"], 2)
        with ZipFile(processed) as archive:
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            rows = sheet.findall(f"{{{NS}}}sheetData/{{{NS}}}row")
            self.assertEqual(len(rows), 2)
            headers = " ".join(rows[0].itertext())
            self.assertNotIn("wtc_InvUpTmp_mean", headers)
            self.assertIn("wtc_GearPres_mean", headers)
            self.assertEqual(len(rows[0]), 7)
        workbook = load_workbook(processed, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Station_2304510", "Station_2304511"])
            self.assertEqual(len(list(workbook.worksheets[0].values)), 2)
        finally:
            workbook.close()

    @unittest.skipUnless(os.name == "nt", "Windows file permission regression")
    def test_published_file_inherits_destination_permissions(self) -> None:
        with tempfile.TemporaryDirectory(dir=self.root) as destination:
            folder = Path(destination)
            # Reset the destination itself, which otherwise also has a private temp ACL.
            subprocess.run(["icacls", str(folder), "/reset"], check=True, capture_output=True)
            expected = folder / "ordinary.txt"
            expected.write_text("test")
            with tempfile.TemporaryDirectory(dir=folder) as private:
                staged = Path(private) / "staged.xlsx"
                staged.write_bytes(b"workbook bytes")
                output = folder / "published.xlsx"
                publish_workbook(staged, output)
            def permissions(path: Path) -> str:
                result = subprocess.run(["icacls", str(path)], check=True,
                                        capture_output=True, text=True)
                return " ".join(result.stdout.replace(str(path), "").split())
            self.assertEqual(permissions(output), permissions(expected))
            self.assertEqual(output.read_bytes(), b"workbook bytes")

    def test_identical_sensor_duplicates_are_reported(self) -> None:
        row = ["2018-06-01 00:00:00", "2304510", "1"]
        self.write("tblSCTurPress", ["TimeStamp", "StationId", "wtc_GearPres_mean"],
                   [row, row])
        report = export_turbines(self.root, self.output)
        self.assertEqual(report["exact_duplicates_removed"]["tblSCTurPress"], 1)

    def test_cleaning_removes_partially_missing_sensor_rows(self) -> None:
        self.write("tblSCTurbine", ["TimeStamp", "StationId", "wtc_PrWindSp_mean"], [
            ["2018-06-01 00:00:00", "2304510", "1.5"],
            ["2018-06-01 00:00:00", "2304511", ""],
        ])
        export_turbines(self.root, self.output)
        report = clean_workbook(self.output, self.root / "processed.xlsx")
        self.assertEqual(report["sheets"]["Station_2304511"]["kept_rows"], 0)
        self.assertEqual(report["sheets"]["Station_2304511"]["removed_rows"], 1)


if __name__ == "__main__":
    unittest.main()
