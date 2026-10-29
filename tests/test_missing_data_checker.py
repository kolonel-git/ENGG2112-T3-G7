"""Validate missing-data detection independently of the extraction pipeline."""

from openpyxl import Workbook

from check_excel_missing_data import check_workbook


def test_missing_values_errors_formulas_and_absent_headers(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["TimeStamp", "StationId", "sensor", "Alarmcode"])
    sheet.append(["2018-06-01", "1", 0, None])
    sheet.append(["2018-06-02", "1", None, None])
    sheet.append(["2018-06-03", "1", " NaN ", None])
    sheet.append(["2018-06-04", "1", "#N/A", None])
    sheet.append(["2018-06-05", "1", "=1+1", None])
    workbook.save(path)
    summary, details = check_workbook(path, ["TimeStamp", "StationId", "sensor", "absent"])
    assert summary[0]["status"] == "FAIL"
    assert summary[0]["rows_with_issues"] == 4
    assert summary[0]["missing_columns"] == "absent"
    sensor = next(item for item in details if item["feature"] == "sensor")
    assert sensor["missing_cells"] == 2
    assert sensor["error_cells"] == 1
    assert sensor["formula_cells"] == 1
    assert sensor["example_excel_rows"] == "3; 4; 5; 6"
    assert not any(item["feature"] == "Alarmcode" for item in details)


def test_complete_sheet_and_empty_sheet(tmp_path):
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    workbook.active.append(["sensor"])
    workbook.active.append([0])
    workbook.create_sheet("empty").append(["sensor"])
    workbook.save(path)
    summary, _ = check_workbook(path, ["sensor"])
    assert summary[0]["status"] == "PASS"
    assert summary[1]["status"] == "FAIL"
    assert summary[1]["error"] == "Empty sheet"
