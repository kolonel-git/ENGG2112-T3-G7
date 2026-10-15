"""Read-only batch check of missing cells in combined turbine Excel workbooks.

Run: python tests/check_excel_missing_data.py
Or:  python tests/check_excel_missing_data.py --path data/processed
Exit codes: 0 = complete; 1 = missing data/schema issues; 2 = unreadable input.
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from wt_predmaint.data.turbine_export import SENSOR_COLUMNS  # noqa: E402

# CHANGE THIS to one .xlsx file OR a folder (subfolders are scanned automatically).
CHECK_PATH = ROOT / "data" / "processed"

# Alarm blanks mean no recorded alarm start, so they are not missing sensor readings.
IGNORE_COLUMNS = {"Alarmcode", "AlarmEvents"}

# These two features were unavailable in the June source data. Their absence is allowed.
# If present in a workbook, their cells ARE checked like any other feature.
ALLOWED_ABSENT_COLUMNS = {"wtc_InvUpTmp_mean", "wtc_TrafoTMa_mean"}
REQUIRED_COLUMNS = ["TimeStamp", "StationId", *(
    name for name in SENSOR_COLUMNS if name not in ALLOWED_ABSENT_COLUMNS
)]

# This is a count of example Excel row numbers per feature, not a cap on checking.
MAX_EXAMPLE_ROWS = 10
REPORT_DIR = ROOT / "data" / "processed" / "missing_data_checks"
MISSING_TEXT = {"", "nan", "na", "n/a", "null", "none", "nat", "<na>"}


def is_missing(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in MISSING_TEXT
    return isinstance(value, float) and not math.isfinite(value)


def check_workbook(path: Path, required_columns=None) -> tuple[list[dict], list[dict]]:
    """Return one summary per sheet and one detail per checked column.

    Excel errors and formulas are flagged too: exported readings should be literal
    values, and an uncached formula must not silently pass a completeness check.
    This checks cell completeness, not gaps in the expected sampling timeline.
    """
    required = REQUIRED_COLUMNS if required_columns is None else required_columns
    summaries, details = [], []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            # Read the actual worksheet XML, even when its dimension metadata is wrong.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            header = next(iterator, ())
            names = [str(c.value).strip() if c.value is not None else "" for c in header]
            absent = sorted(set(required) - set(names))
            duplicates = sorted({name for name in names if name and names.count(name) > 1})
            checked = {i: name or f"<unnamed column {i + 1}>"
                       for i, name in enumerate(names) if name not in IGNORE_COLUMNS}
            counts = {i: {"missing_cells": 0, "error_cells": 0, "formula_cells": 0,
                          "examples": []} for i in checked}
            row_count, bad_rows = 0, 0
            for row_number, row in enumerate(iterator, 2):
                row_count += 1
                bad = False
                for index in checked:
                    cell = row[index] if index < len(row) else None
                    kind = cell.data_type if cell is not None else None
                    missing = is_missing(cell.value if cell is not None else None)
                    stats = counts[index]
                    stats["missing_cells"] += int(missing)
                    stats["error_cells"] += int(kind == "e")
                    stats["formula_cells"] += int(kind == "f")
                    if missing or kind in {"e", "f"}:
                        bad = True
                        if len(stats["examples"]) < MAX_EXAMPLE_ROWS:
                            stats["examples"].append(row_number)
                bad_rows += int(bad)
            status = "FAIL" if (bad_rows or absent or duplicates or not row_count
                                  or "" in names) else "PASS"
            summary = {"file": str(path.resolve()), "sheet": sheet.title, "status": status,
                       "rows_checked": row_count, "rows_with_issues": bad_rows,
                       "missing_cells": sum(c["missing_cells"] for c in counts.values()),
                       "missing_columns": "; ".join(absent),
                       "duplicate_headers": "; ".join(duplicates),
                       "error": "Empty sheet" if not row_count else ""}
            summaries.append(summary)
            for index, name in checked.items():
                stats = counts[index]
                details.append({"file": str(path.resolve()), "sheet": sheet.title,
                                "feature": name, "rows_checked": row_count,
                                "missing_cells": stats["missing_cells"],
                                "missing_percent": round(100 * stats["missing_cells"] /
                                                         row_count, 4) if row_count else 0,
                                "error_cells": stats["error_cells"],
                                "formula_cells": stats["formula_cells"],
                                "example_excel_rows": "; ".join(map(str, stats["examples"]))})
    finally:
        workbook.close()
    return summaries, details


def write_csv(path: Path, records: list[dict], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--path", type=Path, default=CHECK_PATH)
    parser.add_argument("--report-dir", type=Path, default=REPORT_DIR)
    args = parser.parse_args()
    files = [args.path] if args.path.is_file() else sorted(args.path.rglob("*.xlsx"))
    files = [p for p in files if p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")]
    if not files:
        print(f"No Excel workbooks found: {args.path}", file=sys.stderr)
        return 2
    summaries, details = [], []
    exit_code = 0
    for path in files:
        print(f"Checking {path} ...", flush=True)
        try:
            sheet_summaries, feature_details = check_workbook(path)
            summaries.extend(sheet_summaries)
            details.extend(feature_details)
            for item in sheet_summaries:
                print(f"  {item['status']} {item['sheet']}: "
                      f"{item['rows_with_issues']:,}/{item['rows_checked']:,} rows with issues; "
                      f"{item['missing_cells']:,} missing cells", flush=True)
                if item["status"] != "PASS":
                    exit_code = max(exit_code, 1)
                    if item["missing_columns"]:
                        print(f"    Missing columns: {item['missing_columns']}", flush=True)
        except Exception as exc:
            # Continue checking the remaining workbooks when one is unreadable.
            summaries.append({"file": str(path.resolve()), "sheet": "", "status": "ERROR",
                              "error": f"{type(exc).__name__}: {exc}"})
            print(f"  ERROR: {exc}", file=sys.stderr)
            exit_code = 2
    args.report_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.report_dir / "summary.csv", summaries, [
        "file", "sheet", "status", "rows_checked", "rows_with_issues", "missing_cells",
        "missing_columns", "duplicate_headers", "error",
    ])
    write_csv(args.report_dir / "feature_details.csv", details, [
        "file", "sheet", "feature", "rows_checked", "missing_cells", "missing_percent",
        "error_cells", "formula_cells", "example_excel_rows",
    ])
    print(f"Reports: {args.report_dir.resolve()}")
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
